import os
import sys
import json
import subprocess
import argparse
from tkinter import Tk, filedialog
import shutil
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage



def ask_file(title="Select manifest JSON file"):
    root = Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title=title,
        filetypes=[("JSON files", "*.json"), ("All files", "*.*")]
    )
    root.destroy()
    return path


def format_excel_results(category, results_dir, template_path):
    """
    Populates the PET_Results_Template.xlsx based on CSV and PNG outputs.
    Saves as PET_Results_<category>.xlsx inside results_dir.
    """
    csv_files = [f for f in os.listdir(results_dir) if f.lower().endswith(".csv")]
    png_files = [f for f in os.listdir(results_dir) if f.lower().endswith(".png")]
    if not csv_files:
        print(f"⚠️ No CSV found for {category}. Skipping Excel formatting.")
        return

    csv_path = os.path.join(results_dir, csv_files[0])
    df = pd.read_csv(csv_path)

    out_xlsx = os.path.join(results_dir, f"PET_Results_{category}.xlsx")

    try:
        wb = load_workbook(template_path)
        ws = wb.active

        if category == "acr":
            # Background SUVs
            bg_vals = df["SUVmean"].values[:3] if len(df) >= 3 else [None]*3
            for i, v in enumerate(bg_vals, start=3):
                ws[f"C{i}"] = float(v) if v is not None else ""

            # Max hot cell SUVs 25,16,12,8 mm
            hot_vals = df["SUVmax"].values[:4] if len(df) >= 4 else [None]*4
            for i, v in enumerate(hot_vals, start=4):
                ws[f"G{i}"] = float(v) if v is not None else ""

            # Bone, air, water mean/min
            try:
                ws["G9"] = float(df.loc[df["Region"]=="Bone","SUVmean"].values[0])
                ws["H9"] = float(df.loc[df["Region"]=="Bone","SUVmin"].values[0])
                ws["G10"] = float(df.loc[df["Region"]=="Air","SUVmean"].values[0])
                ws["H10"] = float(df.loc[df["Region"]=="Air","SUVmin"].values[0])
                ws["G11"] = float(df.loc[df["Region"]=="Water","SUVmean"].values[0])
                ws["H11"] = float(df.loc[df["Region"]=="Water","SUVmin"].values[0])
            except Exception:
                pass

        elif category == "crp":
            # Count Rate Performance SUVs
            if "SUVmean" in df.columns:
                crp_vals = df["SUVmean"].values[:3] if len(df) >= 3 else [None]*3
                for i, v in enumerate(crp_vals, start=3):
                    ws[f"K{i}"] = float(v) if v is not None else ""

        elif category == "uniformity":
            # Uniformity might not have scalar data, just images
            pass

        # --- Insert images if present ---
        def safe_add_img(anchor, partial_name):
            for f in png_files:
                if partial_name.lower() in f.lower():
                    try:
                        img = XLImage(os.path.join(results_dir, f))
                        ws.add_image(img, anchor)
                        print(f"🖼️ Added {f} at {anchor}")
                        return
                    except Exception as e:
                        print(f"⚠️ Could not add {f}: {e}")

        safe_add_img("B18", "2x2")
        safe_add_img("F18", "hot")
        safe_add_img("J18", "count")
        safe_add_img("B29", "uniformity")
        safe_add_img("F29", "roi")

        wb.save(out_xlsx)
        print(f"✅ Excel formatted: {out_xlsx}")

    except Exception as e:
        print(f"⚠️ Excel formatting failed for {category}: {e}")


def run_analysis_scripts_from_manifest(manifest_path, out_root=None):
    # Load manifest
    with open(manifest_path, "r", encoding="utf-8") as f:
        manifest = json.load(f)

    # If output not specified, use parent of manifest
    if out_root is None:
        out_root = os.path.dirname(os.path.abspath(manifest_path))

    python_results_root = os.path.join(out_root, "Python_Results")
    os.makedirs(python_results_root, exist_ok=True)

    # Map manifest category → analysis script
    script_map = {
        "acr": "3b._Auto_SUV_Overlay.py",
        "uniformity": "4b._Auto_Uniformity_Scoring.py",
        "crp": "5._Count_Rate_Performance.py"
    }

    print(f"\n📄 Loaded manifest: {manifest_path}")
    print(f"📁 Output root: {python_results_root}\n")

    # Iterate over each category in manifest
    for label, entry in manifest.items():
        if label not in script_map:
            print(f"⚠️ Skipping '{label}' — no associated script.")
            continue

        # ----------------------------------------------------
        # Support BOTH old-style and new-style manifest shapes
        # ----------------------------------------------------
        in_dir = None

        # Old style: list of dicts: { "acr": [ { "rep_file": ... }, ... ] }
        if isinstance(entry, list):
            for item in entry:
                # Prefer rep_file if present
                rep_file = item.get("rep_file")
                if rep_file and os.path.exists(rep_file):
                    in_dir = os.path.dirname(rep_file)
                    break

                # Fallback: folder key if present
                folder = item.get("folder")
                if folder and os.path.isdir(folder):
                    in_dir = folder
                    break

        # New style: single dict: { "acr": { "folder": "...", ... } }
        elif isinstance(entry, dict):
            rep_file = entry.get("rep_file")
            folder = entry.get("folder")

            if rep_file and os.path.exists(rep_file):
                in_dir = os.path.dirname(rep_file)
            elif folder and os.path.isdir(folder):
                in_dir = folder

        if not in_dir:
            print(f"⚠️ Skipping '{label}' — no valid folder/rep_file found.")
            continue

        # --- Input/output directories ---
        out_dir = os.path.join(python_results_root, label)
        os.makedirs(out_dir, exist_ok=True)

        script_path = os.path.join(os.path.dirname(__file__), script_map[label])
        if not os.path.exists(script_path):
            print(f"⚠️ Script not found for '{label}': {script_path}")
            continue

        # --- Run the analysis script ---
        env = os.environ.copy()
        env["PYTHONIOENCODING"] = "utf-8"

        cmd = [
            sys.executable,
            script_path,
            "--input", in_dir,
            "--output", out_dir
        ]

        print(f"\n🚀 Running {os.path.basename(script_path)} for '{label}'...")
        print(f"   Input : {in_dir}")
        print(f"   Output: {out_dir}")

        try:
            subprocess.run(
                cmd,
                check=True,
                cwd=os.path.dirname(script_path),
                env=env
            )
            print(f"✅ Completed '{label}'")
        except subprocess.CalledProcessError as e:
            print(f"❌ Failed '{label}': {e}")
            continue

        # --- After success, copy .png and .csv files into Python_Results root ---
        for root, _, files in os.walk(out_dir):
            for fn in files:
                if fn.lower().endswith((".png", ".csv")):
                    src = os.path.join(root, fn)
                    dst = os.path.join(python_results_root, fn)

                    # Prefix category if file already exists or is ambiguous
                    if os.path.exists(dst):
                        base, ext = os.path.splitext(fn)
                        dst = os.path.join(python_results_root, f"{label}_{base}{ext}")

                    try:
                        shutil.copy2(src, dst)
                        print(f"📄 Copied {fn} → {os.path.basename(dst)}")
                    except Exception as e:
                        print(f"⚠️ Could not copy {fn}: {e}")

        template_path = os.path.join(os.path.dirname(__file__), "PET_Results_Template.xlsx")
        if os.path.exists(template_path):
            format_excel_results(label, python_results_root, template_path)
        else:
            print(f"⚠️ Template not found: {template_path}")

        # --- Optionally remove the category subfolder entirely ---
        try:
            shutil.rmtree(out_dir)
        except Exception as e:
            print(f"⚠️ Could not remove temp folder {out_dir}: {e}")


def main():
    parser = argparse.ArgumentParser(
        description="Run PET phantom analysis scripts from a manifest JSON."
    )
    parser.add_argument(
        "--manifest", "-m",
        help="Path to the manifest JSON (e.g., series_manifest.json)"
    )
    parser.add_argument(
        "--output", "-o",
        help="Optional root output folder (default: manifest's parent folder)"
    )
    args = parser.parse_args()

    # ✅ File picker if manifest not provided
    manifest_path = args.manifest or ask_file("Select the manifest JSON (series_manifest.json)")
    if not manifest_path:
        print("❌ No manifest file selected. Exiting.")
        return

    run_analysis_scripts_from_manifest(manifest_path, args.output)


if __name__ == "__main__":
    main()
