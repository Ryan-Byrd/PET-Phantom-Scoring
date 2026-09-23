import os
import sys
import json
import subprocess
import argparse
import importlib
import traceback
from tkinter import Tk, filedialog, messagebox
import shutil
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment


def get_repo_root():
    """Get the absolute path to the repository root."""
    # This file is in pet_tools/, so go up one level
    here = os.path.dirname(os.path.abspath(__file__))
    repo = os.path.dirname(here)
    return repo


def ask_file(title="Select manifest JSON file"):
    root = Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title=title,
        filetypes=[("JSON files", "*.json"), ("All files", "*.*")]
    )
    root.destroy()
    return path


def format_excel_results(category, results_dir, template_path):
    """
    Populates the PET_Results_Template.xlsx based on CSV and PNG outputs.
    Saves as PET_Results_<category>.xlsx inside results_dir.
    """
    csv_files = [f for f in os.listdir(results_dir) if f.lower().endswith(".csv")]
    png_files = [f for f in os.listdir(results_dir) if f.lower().endswith(".png")]
    if not csv_files:
        print(f"⚠️ No CSV found for {category}. Skipping Excel formatting.")
        return

    csv_path = os.path.join(results_dir, csv_files[0])
    df = pd.read_csv(csv_path)

    out_xlsx = os.path.join(results_dir, f"PET_Results_{category}.xlsx")

    try:
        wb = load_workbook(template_path)
        ws = wb.active

        if category == "acr":
            # Background SUVs
            bg_vals = df["SUVmean"].values[:3] if len(df) >= 3 else [None]*3
            for i, v in enumerate(bg_vals, start=3):
                ws[f"C{i}"] = float(v) if v is not None else ""

            # Max hot cell SUVs 25,16,12,8 mm
            hot_vals = df["SUVmax"].values[:4] if len(df) >= 4 else [None]*4
            for i, v in enumerate(hot_vals, start=4):
                ws[f"G{i}"] = float(v) if v is not None else ""

            # Bone, air, water mean/min
            try:
                ws["G9"] = float(df.loc[df["Region"]=="Bone","SUVmean"].values[0])
                ws["H9"] = float(df.loc[df["Region"]=="Bone","SUVmin"].values[0])
                ws["G10"] = float(df.loc[df["Region"]=="Air","SUVmean"].values[0])
                ws["H10"] = float(df.loc[df["Region"]=="Air","SUVmin"].values[0])
                ws["G11"] = float(df.loc[df["Region"]=="Water","SUVmean"].values[0])
                ws["H11"] = float(df.loc[df["Region"]=="Water","SUVmin"].values[0])
            except Exception:
                pass

        elif category == "crp":
            # Count Rate Performance SUVs
            if "SUVmean" in df.columns:
                crp_vals = df["SUVmean"].values[:3] if len(df) >= 3 else [None]*3
                for i, v in enumerate(crp_vals, start=3):
                    ws[f"K{i}"] = float(v) if v is not None else ""

        elif category == "uniformity":
            # Uniformity might not have scalar data, just images
            pass

        # --- Insert images if present ---
        def safe_add_img(anchor, partial_name):
            for f in png_files:
                if partial_name.lower() in f.lower():
                    try:
                        img = XLImage(os.path.join(results_dir, f))
                        ws.add_image(img, anchor)
                        print(f"🖼️ Added {f} at {anchor}")
                        return
                    except Exception as e:
                        print(f"⚠️ Could not add {f}: {e}")

        safe_add_img("B18", "2x2")
        safe_add_img("F18", "hot")
        safe_add_img("J18", "count")
        safe_add_img("B29", "uniformity")
        safe_add_img("F29", "roi")

        wb.save(out_xlsx)
        print(f"✅ Excel formatted: {out_xlsx}")

    except Exception as e:
        print(f"⚠️ Excel formatting failed for {category}: {e}")


def combine_summary_csvs_to_excel(results_root, log):
    """
    Combine key output CSV files into one labeled Excel workbook.

    The workbook contains one sheet where each CSV is written as a separate
    section with a bold title row. Section spacing is explicit so report
    layout can be tuned for readability.
    """
    def safe_log(message):
        try:
            log(message)
        except Exception:
            # Logging should never prevent workbook generation.
            pass

    csv_sections = [
        ("Count Rate Performance", "Count_Rate_Performance_SUV_Measurements.csv"),
        ("ACR SUV Results", "SUV_results.csv"),
    ]

    present_sections = []
    for section_title, filename in csv_sections:
        csv_path = os.path.join(results_root, filename)
        if os.path.isfile(csv_path):
            present_sections.append((section_title, csv_path))

    if not present_sections:
        safe_log("WARNING: No summary CSV files found to combine into Excel.")
        return None

    out_xlsx = os.path.join(results_root, "Combined_Analysis_Results.xlsx")

    try:
        with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
            sheet_name = "Combined Results"
            start_row = 0
            title_rows = []

            for section_title, csv_path in present_sections:
                try:
                    df = pd.read_csv(csv_path)
                except Exception as e:
                    safe_log(f"WARNING: Skipping '{os.path.basename(csv_path)}' due to read error: {e}")
                    continue

                # Section title
                pd.DataFrame({"Section": [section_title]}).to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                    header=False,
                    startrow=start_row,
                    startcol=0,
                )
                title_rows.append(start_row + 1)  # openpyxl uses 1-based rows

                # Section table
                safe_df = df.where(pd.notnull(df), "")

                if section_title == "ACR SUV Results" and not safe_df.empty:
                    first_col = safe_df.columns[0]
                    marker_mask = safe_df[first_col].astype(str).str.contains(
                        r"uniformity\s+statistics",
                        case=False,
                        regex=True,
                    )
                    marker_indices = safe_df.index[marker_mask]
                    if len(marker_indices) > 0:
                        marker_pos = safe_df.index.get_loc(marker_indices[0])
                        if marker_pos > 0:
                            blank_row = pd.DataFrame([{col: "" for col in safe_df.columns}])
                            safe_df = pd.concat(
                                [
                                    safe_df.iloc[:marker_pos],
                                    blank_row,
                                    safe_df.iloc[marker_pos:],
                                ],
                                ignore_index=True,
                            )

                safe_df.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                    header=True,
                    startrow=start_row + 1,
                    startcol=0,
                )

                # Add one blank separator row between top-level sections.
                start_row += len(safe_df.index) + 3

            ws = writer.book[sheet_name]
            for row_idx in title_rows:
                ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
                ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")

            for row_idx in range(1, ws.max_row + 1):
                cell_value = ws.cell(row=row_idx, column=1).value
                if isinstance(cell_value, str) and cell_value in [s[0] for s in csv_sections]:
                    ws.cell(row=row_idx, column=1).font = Font(bold=True)

        safe_log(f"OK: Combined Excel created: {out_xlsx}")
        return out_xlsx
    except Exception as e:
        safe_log(f"WARNING: Failed to create combined Excel workbook: {e}")
        return None


def consolidate_csv_files(results_root, log):
    """
    Move all CSV files under results_root into a single folder.

    This keeps final outputs centralized in one location.
    """
    def safe_log(message):
        try:
            log(message)
        except Exception:
            pass

    csv_bundle_dir = os.path.join(results_root, "All_CSV_Files")
    os.makedirs(csv_bundle_dir, exist_ok=True)

    copied_count = 0
    for root, _, files in os.walk(results_root):
        # Prevent recursively copying files already placed in the bundle folder.
        if os.path.abspath(root).startswith(os.path.abspath(csv_bundle_dir)):
            continue

        for filename in files:
            if not filename.lower().endswith(".csv"):
                continue

            src = os.path.join(root, filename)
            rel_dir = os.path.relpath(root, results_root)
            rel_tag = "root" if rel_dir == "." else rel_dir.replace(os.sep, "_")

            dst = os.path.join(csv_bundle_dir, filename)
            if os.path.exists(dst):
                base, ext = os.path.splitext(filename)
                dst = os.path.join(csv_bundle_dir, f"{rel_tag}_{base}{ext}")

            try:
                shutil.move(src, dst)
                copied_count += 1
            except Exception as e:
                safe_log(f"WARNING: Could not move CSV '{src}' to bundle folder: {e}")

    safe_log(f"OK: Moved {copied_count} CSV file(s) into: {csv_bundle_dir}")
    return csv_bundle_dir

def _run_module_main(module_name, argv):
    """
    Execute module.main(argv) and normalize return codes.

    This avoids filesystem script-path assumptions and works in both source
    and PyInstaller-frozen environments.
    """
    module = importlib.import_module(module_name)
    if not hasattr(module, "main"):
        raise RuntimeError(f"Module '{module_name}' does not expose a main(argv) entry point.")

    try:
        result = module.main(argv)
    except SystemExit as e:
        code = getattr(e, "code", 0)
        if isinstance(code, int):
            return code
        return 0

    if result is None:
        return 0
    if isinstance(result, int):
        return result
    return 0


def run_analysis_scripts_from_manifest(manifest_path, out_root=None, debug=False):
    """
    Run analysis scripts based on manifest JSON.
    
    Args:
        manifest_path: Path to series_manifest.json
        out_root: Optional output root (default: manifest's parent directory)
    """
    # Ensure absolute path for manifest
    manifest_path = os.path.abspath(manifest_path)
    
    if not os.path.exists(manifest_path):
        raise FileNotFoundError(f"Manifest not found: {manifest_path}")
    
    # Load manifest
    try:
        with open(manifest_path, "r", encoding="utf-8") as f:
            manifest = json.load(f)
    except json.JSONDecodeError as e:
        raise ValueError(f"Invalid JSON in manifest: {e}")

    # If output not specified, use parent of manifest
    if out_root is None:
        out_root = os.path.dirname(manifest_path)
    
    out_root = os.path.abspath(out_root)

    python_results_root = os.path.join(out_root, "Python_Results")
    os.makedirs(python_results_root, exist_ok=True)

    # Map manifest category → analysis module
    module_map = {
        "acr": "pet_tools.suv_overlay_auto",
        "uniformity": "pet_tools.uniformity_scoring_auto",
        "crp": "pet_tools.count_rate_performance"
    }

    # Get absolute paths for modules
    repo_root = get_repo_root()
    pet_tools_dir = os.path.join(repo_root, "pet_tools")
    ct_toolkit_path = os.path.join(repo_root, "CT_Analysis_Toolkit")
    
    # Ensure paths exist
    if not os.path.isdir(pet_tools_dir):
        raise RuntimeError(f"pet_tools directory not found: {pet_tools_dir}")

    extra_paths = [ct_toolkit_path, repo_root, pet_tools_dir]

    debug_log_path = os.path.join(python_results_root, "run_from_manifest_debug.log") if debug else None

    def log(msg):
        print(msg)
        if debug_log_path:
            try:
                with open(debug_log_path, "a", encoding="utf-8") as f:
                    f.write(msg + "\n")
            except Exception:
                pass

    if debug_log_path:
        try:
            with open(debug_log_path, "w", encoding="utf-8") as f:
                f.write("run_from_manifest debug log\n")
                f.write(f"sys.executable: {sys.executable}\n")
                f.write(f"cwd: {os.getcwd()}\n")
                f.write(f"frozen: {bool(getattr(sys, 'frozen', False))}\n")
        except Exception:
            pass

    log(f"\n📄 Loaded manifest: {manifest_path}")
    log(f"📁 Output root: {python_results_root}")
    log(f"📦 Repo root: {repo_root}\n")

    # Make extra paths importable in source runs; frozen runs generally do not need this.
    for p in extra_paths:
        if os.path.isdir(p) and p not in sys.path:
            sys.path.insert(0, p)

    failures = []

    # Iterate over each category in manifest
    for label, entry in manifest.items():
        if label not in module_map:
            log(f"⚠️  Skipping '{label}' — no associated analysis module.")
            continue

        # Support BOTH old-style and new-style manifest shapes
        in_dir = None

        # Old style: list of dicts: { "acr": [ { "rep_file": ... }, ... ] }
        if isinstance(entry, list):
            for item in entry:
                # Prefer rep_file if present
                rep_file = item.get("rep_file")
                if rep_file:
                    rep_file = os.path.abspath(rep_file)
                    if os.path.exists(rep_file):
                        in_dir = os.path.dirname(rep_file)
                        break

                # Fallback: folder key if present
                folder = item.get("folder")
                if folder:
                    folder = os.path.abspath(folder)
                    if os.path.isdir(folder):
                        in_dir = folder
                        break

        # New style: single dict: { "acr": { "folder": "...", ... } }
        elif isinstance(entry, dict):
            rep_file = entry.get("rep_file")
            folder = entry.get("folder")

            if rep_file:
                rep_file = os.path.abspath(rep_file)
                if os.path.exists(rep_file):
                    in_dir = os.path.dirname(rep_file)
            elif folder:
                folder = os.path.abspath(folder)
                if os.path.isdir(folder):
                    in_dir = folder

        if not in_dir:
            log(f"⚠️  Skipping '{label}' — no valid folder/rep_file found.")
            if isinstance(entry, dict):
                log(f"    Manifest entry: {entry}")
            continue

        # --- Input/output directories ---
        out_dir = os.path.join(python_results_root, label)
        os.makedirs(out_dir, exist_ok=True)

        module_name = module_map[label]

        # --- Run the analysis script ---
        log(f"\n🚀 Running {module_name} for '{label}'...")
        log(f"   Input : {in_dir}")
        log(f"   Output: {out_dir}")
        if debug:
            log(f"   sys.path head: {sys.path[:5]}")

        try:
            result_code = _run_module_main(
                module_name,
                ["--input", in_dir, "--output", out_dir],
            )
            
            if result_code == 0:
                log(f"✅ Completed '{label}'")
                if label == "acr":
                    # Representative-slice export is a separate tool:
                    # acr_slice_selector identifies the uniformity/rods/hot-cell
                    # center slices, then calls image_selection headlessly
                    # (--indices <3 slices> --fov 220 --save). suv_overlay_auto
                    # does not do this, so run it explicitly here.
                    log("ℹ️  Running acr_slice_selector to export representative slices...")
                    try:
                        selector_args = ["--input", in_dir, "--output", out_dir]
                        if debug:
                            selector_args.append("--debug")
                        sel_code = _run_module_main(
                            "pet_tools.acr_slice_selector",
                            selector_args,
                        )
                        if sel_code == 0:
                            log("✅ Representative slices exported via image_selection")
                        else:
                            log(f"⚠️  acr_slice_selector returned code {sel_code}; continuing")
                    except Exception as e:
                        log(f"⚠️  acr_slice_selector failed: {e}; continuing")
            else:
                log(f"❌ Failed '{label}' with return code {result_code}")
                failures.append((label, f"return code {result_code}"))
                continue
                
        except Exception as e:
            log(f"❌ Exception running '{label}': {e}")
            if debug:
                log(traceback.format_exc())
            failures.append((label, str(e)))
            continue

        # --- After success, copy .png and .csv files into Python_Results root ---
        for root, _, files in os.walk(out_dir):
            for fn in files:
                if fn.lower().endswith((".png", ".csv")):
                    src = os.path.join(root, fn)
                    dst = os.path.join(python_results_root, fn)

                    # Prefix category if file already exists or is ambiguous
                    if os.path.exists(dst):
                        base, ext = os.path.splitext(fn)
                        dst = os.path.join(python_results_root, f"{label}_{base}{ext}")

                    try:
                        shutil.copy2(src, dst)
                        log(f"📄 Copied {fn} → {os.path.basename(dst)}")
                    except Exception as e:
                        log(f"⚠️  Could not copy {fn}: {e}")

        # --- Optionally remove the category subfolder entirely ---
        try:
            shutil.rmtree(out_dir)
        except Exception as e:
            log(f"⚠️  Could not remove temp folder {out_dir}: {e}")

    if failures:
        log(f"\n⚠️  Completed with failures: {len(failures)}")
        for label, reason in failures:
            log(f"   - {label}: {reason}")
    else:
        log(f"\n✨ All analyses completed. Results in: {python_results_root}")

    # Final consolidation step: combine summary CSV outputs into one workbook.
    combine_summary_csvs_to_excel(python_results_root, log)
    consolidate_csv_files(python_results_root, log)

    return failures





def main():
    parser = argparse.ArgumentParser(
        description="Run PET phantom analysis scripts from a manifest JSON."
    )
    parser.add_argument(
        "--manifest", "-m",
        help="Path to the manifest JSON (e.g., series_manifest.json)"
    )
    parser.add_argument(
        "--output", "-o",
        help="Optional root output folder (default: manifest's parent folder)"
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="Enable debug logging and write run_from_manifest_debug.log"
    )
    args = parser.parse_args()

    # File picker if manifest not provided
    manifest_path = args.manifest or ask_file("Select the manifest JSON (series_manifest.json)")
    if not manifest_path:
        print("❌ No manifest file selected. Exiting.")
        return 1

    try:
        failures = run_analysis_scripts_from_manifest(manifest_path, args.output, debug=args.debug)
        if failures:
            print("\n❌ One or more analyses failed. See console output for details.")
            return 1
        return 0
    except Exception as e:
        print(f"\n❌ Error: {e}")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(main())
